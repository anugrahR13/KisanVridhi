import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, title, text, border_color="16A34A", bg_color="F0FDF4"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_title = p.add_run(f"★ {title}\n")
    r_title.bold = True
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(11)
    r_title.font.color.rgb = RGBColor(15, 81, 50)
    
    r_text = p.add_run(text)
    r_text.font.name = "Times New Roman"
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph()

def add_heading_1(doc, text):
    try:
        p = doc.add_paragraph(style='Heading 1')
    except Exception:
        p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(13, 83, 40)
    return p

def add_heading_2(doc, text):
    try:
        p = doc.add_paragraph(style='Heading 2')
    except Exception:
        p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_heading_3(doc, text):
    try:
        p = doc.add_paragraph(style='Heading 3')
    except Exception:
        p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_body_p(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_bullet_p(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_bold = p.add_run(bold_prefix + ": ")
    r_bold.bold = True
    r_bold.font.name = "Times New Roman"
    r_bold.font.size = Pt(11)
    r_bold.font.color.rgb = RGBColor(15, 23, 42)
    r_text = p.add_run(text)
    r_text.font.name = "Times New Roman"
    r_text.font.size = Pt(11)
    r_text.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_custom_table(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    # Header Row
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], "0D5328")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(255, 255, 255)
            
    # Data Rows
    for r_idx, row in enumerate(data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=140, right=140)
            p = row_cells[c_idx].paragraphs[0]
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(30, 41, 59)
                
    # Widths
    if col_widths:
        for row in tbl.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)
                
    doc.add_paragraph()
    return tbl

def add_styled_heading(doc, text, level=1):
    if level == 1:
        return add_heading_1(doc, text)
    elif level == 2:
        return add_heading_2(doc, text)
    else:
        return add_heading_3(doc, text)

def add_body_paragraph(doc, text):
    return add_body_p(doc, text)

def add_bullet_point(doc, text):
    if ":" in text:
        prefix, rest = text.split(":", 1)
        return add_bullet_p(doc, prefix.strip(), rest.strip())
    else:
        return add_bullet_p(doc, "Note", text)

def add_callout_box(doc, title, text):
    return add_callout(doc, title, text)

def create_styled_table(doc, headers, rows, col_widths=None):
    return add_custom_table(doc, headers, rows, col_widths)

def add_page_break(doc):
    doc.add_page_break()

print("Helper functions compiled successfully.")
